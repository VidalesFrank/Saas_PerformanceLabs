"""
Export de resultados del pushover no lineal de pórticos a XLSX.

Hojas producidas:
  1. Resumen          : parámetros + resultados por dirección
  2. Pushover_X       : curva completa paso-a-paso (deriva, cortante basal)
  3. Pushover_Y       : idem
  4. Rotulas          : todas las rótulas con niveles ASCE 41 y DCR
  5. Elementos        : metadata de columnas y vigas

Uso:
    from app.services.frame_pushover_export import build_xlsx_bytes
    data = json.load(open(result_path))
    content = build_xlsx_bytes(data)
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


# Colores por nivel ASCE 41 (hex RGB sin #)
_LEVEL_FILL = {
    "none":     "F3F4F6",
    "near_io":  "DBEAFE",
    "io":       "DCFCE7",
    "ls":       "FEF9C3",
    "cp":       "FFEDD5",
    "collapse": "FEE2E2",
}
_LEVEL_TEXT = {
    "none":     "Elástico",
    "near_io":  "Casi IO",
    "io":       "IO (Ocupación Inmediata)",
    "ls":       "LS (Seguridad de Vida)",
    "cp":       "CP (Prevención de Colapso)",
    "collapse": "Colapso",
}

_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_HEADER_FILL = PatternFill("solid", fgColor="1F2937")
_TITLE_FONT  = Font(bold=True, size=14, color="1F2937")
_SUB_FONT    = Font(italic=True, size=9, color="6B7280")


def _set_header(ws, row: int, cols: list[str]) -> None:
    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=j, value=c)
        cell.font      = _HEADER_FONT
        cell.fill      = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")


def _autofit(ws, max_width: int = 30) -> None:
    for col in ws.columns:
        maxlen = 0
        letter = get_column_letter(col[0].column)
        for cell in col:
            v = cell.value
            if v is None:
                continue
            l = len(str(v))
            if l > maxlen:
                maxlen = l
        ws.column_dimensions[letter].width = min(max(maxlen + 2, 10), max_width)


def _sheet_resumen(wb: Workbook, data: dict[str, Any]) -> None:
    ws = wb.create_sheet("Resumen", 0)
    ws["A1"] = "Pushover No Lineal — Pórticos RC"
    ws["A1"].font = _TITLE_FONT
    ws.merge_cells("A1:D1")

    ws["A2"] = f"Project: {data.get('project_id', '')}"
    ws["A2"].font = _SUB_FONT
    ws["A3"] = f"Computed at: {data.get('computed_at', '')}"
    ws["A3"].font = _SUB_FONT

    row = 5
    ws.cell(row=row, column=1, value="Parámetros").font = Font(bold=True, size=12)
    row += 1
    _set_header(ws, row, ["Parámetro", "Valor"])
    row += 1
    for k, v in [
        ("Patrón de carga",      data.get("pattern_type", "")),
        ("Deriva objetivo (%)",  data.get("target_drift_pct", "")),
        ("Integration points",   data.get("n_integration_pts", "")),
    ]:
        ws.cell(row=row, column=1, value=k)
        ws.cell(row=row, column=2, value=v)
        row += 1

    row += 2
    ws.cell(row=row, column=1, value="Resultados por dirección").font = Font(bold=True, size=12)
    row += 1
    _set_header(ws, row, [
        "Dir", "Status", "Patrón", "Pasos conv.", "Pasos tot.",
        "Deriva máx. (%)", "Vb máx. (kN)", "Rótulas críticas", "DCR máx.",
    ])
    row += 1
    for d, s in (data.get("summary", {}) or {}).items():
        ws.cell(row=row, column=1, value=d)
        ws.cell(row=row, column=2, value=s.get("status"))
        ws.cell(row=row, column=3, value=s.get("pattern_type"))
        ws.cell(row=row, column=4, value=s.get("converged_steps"))
        ws.cell(row=row, column=5, value=s.get("total_steps"))
        ws.cell(row=row, column=6, value=round(s.get("max_drift_pct", 0.0), 4))
        ws.cell(row=row, column=7, value=round(s.get("max_base_shear_kN", 0.0), 2))
        ws.cell(row=row, column=8, value=s.get("n_critical_hinges", 0))
        ws.cell(row=row, column=9, value=round(s.get("max_dcr", 0.0), 3))
        row += 1

    _autofit(ws)


def _sheet_pushover(wb: Workbook, name: str, dir_result: dict[str, Any] | None) -> None:
    ws = wb.create_sheet(name)
    if not dir_result:
        ws["A1"] = "Sin resultados"
        return
    ws["A1"] = f"Pushover — {dir_result.get('direction', name)}"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = (
        f"Patrón: {dir_result.get('pattern_type', '')} · "
        f"Status: {dir_result.get('status', '')} · "
        f"Pasos: {dir_result.get('converged_steps', 0)}/{dir_result.get('total_steps', 0)}"
    )
    ws["A2"].font = _SUB_FONT

    _set_header(ws, 4, ["Step", "Desplazamiento (m)", "Deriva (%)", "Cortante basal (kN)"])
    row = 5
    for s in dir_result.get("steps", []):
        ws.cell(row=row, column=1, value=s.get("step"))
        ws.cell(row=row, column=2, value=s.get("displacement_m"))
        ws.cell(row=row, column=3, value=s.get("drift_pct"))
        ws.cell(row=row, column=4, value=s.get("base_shear_kN"))
        row += 1
    _autofit(ws)


def _sheet_hinges(wb: Workbook, data: dict[str, Any]) -> None:
    ws = wb.create_sheet("Rotulas")
    ws["A1"] = "Rótulas plásticas — ASCE 41-17"
    ws["A1"].font = _TITLE_FONT
    ws["A2"] = (
        "Niveles: IO (Ocupación Inmediata) · LS (Seguridad de Vida) · "
        "CP (Prevención de Colapso) · Colapso. DCR = θ_p / θ_CP."
    )
    ws["A2"].font = _SUB_FONT

    _set_header(ws, 4, [
        "Dir", "Tipo", "Frame", "Ext", "Piso",
        "θ_p (mrad)", "θ_y (mrad)",
        "θ_IO (mrad)", "θ_LS (mrad)", "θ_CP (mrad)",
        "DCR", "Nivel", "Paso IO", "Paso LS", "Paso CP",
    ])
    row = 5

    for dir_label in ("X", "Y"):
        dir_result = data.get(f"pushover_{dir_label}")
        if not dir_result:
            continue
        for h in (dir_result.get("damage") or {}).get("hinges", []):
            level = h.get("damage_level", "none")
            fill  = _LEVEL_FILL.get(level, "FFFFFF")
            vals = [
                dir_label,
                "Columna" if h.get("kind") == "column" else "Viga",
                h.get("fid"),
                h.get("end"),
                h.get("story"),
                round((h.get("theta_p_max", 0.0) or 0.0) * 1000, 2),
                round((h.get("theta_y", 0.0) or 0.0) * 1000, 2),
                round((h.get("theta_IO", 0.0) or 0.0) * 1000, 2),
                round((h.get("theta_LS", 0.0) or 0.0) * 1000, 2),
                round((h.get("theta_CP", 0.0) or 0.0) * 1000, 2),
                round(h.get("dcr", 0.0) or 0.0, 3),
                _LEVEL_TEXT.get(level, level),
                h.get("step_first_IO"),
                h.get("step_first_LS"),
                h.get("step_first_CP"),
            ]
            for j, v in enumerate(vals, start=1):
                cell = ws.cell(row=row, column=j, value=v)
                if fill and j == 12:
                    cell.fill = PatternFill("solid", fgColor=fill)
                    cell.font = Font(bold=True)
            row += 1

    _autofit(ws)
    ws.freeze_panes = "A5"


def _sheet_elements(wb: Workbook, data: dict[str, Any]) -> None:
    """Metadata de columnas/vigas (tomado de element_lines de la primera dir disponible)."""
    ws = wb.create_sheet("Elementos")
    ws["A1"] = "Metadata de elementos del modelo no lineal"
    ws["A1"].font = _TITLE_FONT

    element_lines = []
    for d in ("X", "Y"):
        dr = data.get(f"pushover_{d}")
        if dr and dr.get("element_lines"):
            element_lines = dr["element_lines"]
            break

    _set_header(ws, 3, [
        "Frame", "Tipo", "Piso", "Nodo i", "Nodo j",
        "b (m)", "h (m)", "L (m)", "lp (m)", "ele_tag",
    ])
    row = 4
    for e in element_lines:
        ws.cell(row=row, column=1, value=e.get("fid"))
        ws.cell(row=row, column=2, value="Columna" if e.get("kind") == "column" else "Viga")
        ws.cell(row=row, column=3, value=e.get("story"))
        ws.cell(row=row, column=4, value=e.get("node_i_lbl"))
        ws.cell(row=row, column=5, value=e.get("node_j_lbl"))
        ws.cell(row=row, column=6, value=e.get("b_m"))
        ws.cell(row=row, column=7, value=e.get("h_m"))
        ws.cell(row=row, column=8, value=e.get("L_m"))
        ws.cell(row=row, column=9, value=e.get("lp_m"))
        ws.cell(row=row, column=10, value=e.get("ele_tag"))
        row += 1
    _autofit(ws)
    ws.freeze_panes = "A4"


def build_xlsx_bytes(data: dict[str, Any]) -> bytes:
    wb = Workbook()
    # El worksheet default queda vacío — lo removemos
    default = wb.active
    wb.remove(default)

    _sheet_resumen(wb, data)
    _sheet_pushover(wb, "Pushover_X", data.get("pushover_X"))
    _sheet_pushover(wb, "Pushover_Y", data.get("pushover_Y"))
    _sheet_hinges(wb, data)
    _sheet_elements(wb, data)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
